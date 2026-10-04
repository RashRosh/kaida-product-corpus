"""Deterministic official-source extraction, validation, and crosswalk logic."""

from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook

from .io import file_sha256, read_csv
from .normalization import normalize


RELATIONS = {"EXACT", "BROADER", "NARROWER", "NO_MATCH", "NOT_APPLICABLE"}
ENTITY_RELATIONS = {"EXACT", "BROADER", "NARROWER"}
VERIFICATION_STATUSES = {"CURRENT_VERIFIED", "REVIEW_REQUIRED"}
PROVENANCE_TYPES = {
    "LEGACY_WORKBOOK",
    "CURRENT_OFFICIAL_SOURCE",
    "MANUAL_RECONCILIATION",
}


def load_registry(path: Path, root: Path) -> dict[str, Any]:
    registry = json.loads(path.read_text(encoding="utf-8"))
    sources = registry.get("official_sources", [])
    source_ids = [item.get("official_source_id", "") for item in sources]
    if not source_ids or len(source_ids) != len(set(source_ids)):
        raise ValueError("Official source IDs must be present and unique")
    for source in sources:
        for field in ("official_source_id", "agency", "name", "current_url"):
            if not str(source.get(field, "")).strip():
                raise ValueError(
                    f"{source.get('official_source_id', '<unknown>')}: {field} is required"
                )
        dataset_path = str(source.get("dataset_path", "")).strip()
        if dataset_path:
            dataset = root / dataset_path
            if not dataset.exists():
                raise FileNotFoundError(f"Official dataset is missing: {dataset}")
            expected_hash = str(source.get("dataset_sha256", "")).lower()
            actual_hash = file_sha256(dataset)
            if actual_hash != expected_hash:
                raise ValueError(
                    f"{source['official_source_id']}: dataset SHA-256 mismatch: "
                    f"{actual_hash}; expected {expected_hash}"
                )
    evidence = registry.get("legacy_evidence", {})
    if evidence.get("official_source_id") not in set(source_ids):
        raise ValueError("Legacy evidence references an unknown official source ID")
    evidence_path = root / str(evidence.get("workbook_path", ""))
    if not evidence_path.exists():
        raise FileNotFoundError(f"Legacy evidence workbook is missing: {evidence_path}")
    actual_hash = file_sha256(evidence_path)
    expected_hash = str(evidence.get("workbook_sha256", "")).lower()
    if actual_hash != expected_hash:
        raise ValueError(
            f"Legacy workbook SHA-256 mismatch: {actual_hash}; expected {expected_hash}"
        )
    return registry


def _xls_code(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") and text[:-2].isdigit() else text


def extract_kpved_entities(source: dict[str, Any], root: Path) -> list[dict[str, str]]:
    dataset = root / source["dataset_path"]
    workbook = xlrd.open_workbook(str(dataset))
    sheet = workbook.sheet_by_index(0)
    code_pattern = re.compile(source["code_pattern"])
    entities = []
    seen_codes: set[str] = set()
    for row_number in range(3, sheet.nrows):
        code = _xls_code(sheet.cell_value(row_number, 0))
        name_kk = str(sheet.cell_value(row_number, 1)).strip()
        name_ru = str(sheet.cell_value(row_number, 2)).strip()
        if not code and not name_ru and not name_kk:
            continue
        if not code_pattern.fullmatch(code):
            raise ValueError(
                f"{source['official_source_id']}: malformed code {code!r} "
                f"at XLS row {row_number + 1}"
            )
        if code in seen_codes:
            raise ValueError(
                f"{source['official_source_id']}: duplicate official code {code}"
            )
        if not name_ru:
            raise ValueError(
                f"{source['official_source_id']}:{code}: Russian name is required"
            )
        seen_codes.add(code)
        entities.append(
            {
                "official_source_id": source["official_source_id"],
                "official_entity_code": code,
                "official_entity_name_ru": name_ru,
                "official_entity_name_kk": name_kk,
                "code_level": len(code),
                "source_version": source["version"],
                "source_url": source["current_url"],
                "dataset_sha256": source["dataset_sha256"],
            }
        )
    return sorted(
        entities,
        key=lambda row: (
            len(row["official_entity_code"]),
            row["official_entity_code"],
        ),
    )


def extract_legacy_bns_evidence(
    workbook_path: Path,
    products: dict[str, dict[str, str]],
    official_source_id: str,
    expected_url: str,
) -> list[dict[str, str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Products"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    evidence = []
    for values in rows:
        row = dict(zip(headers, values))
        source_codes = [
            item.strip() for item in str(row.get("source_codes") or "").split(";")
            if item.strip()
        ]
        if "BNS" not in source_codes:
            continue
        source_index = source_codes.index("BNS") + 1
        legacy_url = str(row.get(f"source_url_{source_index}") or "").strip()
        if legacy_url != expected_url:
            raise ValueError(
                f"{row.get('candidate_code')}: BNS URL mismatch: {legacy_url!r}"
            )
        candidate_code = str(row.get("candidate_code") or "").strip()
        evidence.append(
            {
                "legacy_candidate_code": candidate_code,
                "legacy_product_name": str(row.get("name_ru") or "").strip(),
                "current_product_id": candidate_code if candidate_code in products else "",
                "legacy_source_code": "BNS",
                "official_source_id": official_source_id,
                "legacy_source_url": legacy_url,
                "official_entity_code": "",
                "evidence_status": "LEGACY_SOURCE_REFERENCE_ONLY",
                "reason": (
                    "The workbook cites the BNS classifiers landing page but stores "
                    "no entity-level BNS code."
                ),
            }
        )
    return sorted(evidence, key=lambda row: row["legacy_candidate_code"])


def validate_decisions(
    decisions: list[dict[str, str]],
    products: dict[str, dict[str, str]],
    sources: dict[str, dict[str, Any]],
    entities: dict[tuple[str, str], dict[str, str]],
) -> tuple[list[dict[str, str]], int]:
    errors = []
    semantic_seen: dict[tuple[str, ...], dict[str, str]] = {}
    ids: set[str] = set()
    by_product_source: defaultdict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    by_product_source_code: defaultdict[tuple[str, str, str], set[str]] = defaultdict(set)
    exact_by_entity: defaultdict[tuple[str, str], set[str]] = defaultdict(set)

    for line_number, row in enumerate(decisions, start=2):
        mapping_id = row.get("mapping_id", "").strip()
        product_id = row.get("product_id", "").strip()
        source_id = row.get("official_source_id", "").strip()
        code = row.get("official_entity_code", "").strip()
        relation = row.get("relation", "").strip()
        verification = row.get("verification_status", "").strip()
        provenance_type = row.get("provenance_type", "").strip()
        if not mapping_id or mapping_id in ids:
            errors.append(f"line {line_number}: mapping_id must be present and unique")
        ids.add(mapping_id)
        if product_id not in products:
            errors.append(f"{mapping_id}: unknown KAIDA Product ID {product_id!r}")
        if source_id not in sources:
            errors.append(f"{mapping_id}: unknown official source ID {source_id!r}")
        if relation not in RELATIONS:
            errors.append(f"{mapping_id}: unsupported relation {relation!r}")
        if verification not in VERIFICATION_STATUSES:
            errors.append(f"{mapping_id}: unsupported verification status {verification!r}")
        if provenance_type not in PROVENANCE_TYPES:
            errors.append(f"{mapping_id}: unsupported provenance type {provenance_type!r}")
        if not row.get("provenance_reference", "").strip():
            errors.append(f"{mapping_id}: provenance_reference is required")
        if relation in ENTITY_RELATIONS:
            if not code:
                errors.append(f"{mapping_id}: {relation} requires official entity code")
            elif (source_id, code) not in entities:
                errors.append(
                    f"{mapping_id}: official entity does not exist: {source_id}:{code}"
                )
        elif code:
            errors.append(f"{mapping_id}: {relation} must not define official entity code")

        semantic = (
            product_id,
            source_id,
            code,
            relation,
            verification,
            provenance_type,
            row.get("provenance_reference", "").strip(),
            row.get("notes", "").strip(),
        )
        semantic_seen.setdefault(semantic, row)
        by_product_source[(product_id, source_id)].append(row)
        by_product_source_code[(product_id, source_id, code)].add(relation)
        if relation == "EXACT":
            exact_by_entity[(source_id, code)].add(product_id)

    for key, relations in by_product_source_code.items():
        if len(relations) > 1:
            errors.append(f"contradictory relations for {key}: {sorted(relations)}")
    for key, rows in by_product_source.items():
        relations = {row["relation"] for row in rows}
        exact_codes = {
            row["official_entity_code"] for row in rows if row["relation"] == "EXACT"
        }
        if len(exact_codes) > 1:
            errors.append(f"multiple EXACT entities for {key}: {sorted(exact_codes)}")
        if relations.intersection(ENTITY_RELATIONS) and relations.intersection(
            {"NO_MATCH", "NOT_APPLICABLE"}
        ):
            errors.append(f"entity relation contradicts terminal relation for {key}")
    for key, product_ids in exact_by_entity.items():
        if len(product_ids) > 1:
            errors.append(
                f"EXACT official entity maps to multiple Products {key}: {sorted(product_ids)}"
            )
    if errors:
        raise ValueError("Invalid official crosswalk decisions:\n- " + "\n- ".join(errors))

    deduped = sorted(
        semantic_seen.values(),
        key=lambda row: (
            row["product_id"],
            row["official_source_id"],
            row.get("official_entity_code", ""),
            row["relation"],
            row["mapping_id"],
        ),
    )
    return deduped, len(decisions) - len(deduped)


def materialize_crosswalk(
    decisions: list[dict[str, str]],
    products: dict[str, dict[str, str]],
    sources: dict[str, dict[str, Any]],
    entities: dict[tuple[str, str], dict[str, str]],
) -> list[dict[str, str]]:
    rows = []
    for decision in decisions:
        source = sources[decision["official_source_id"]]
        entity = entities.get(
            (decision["official_source_id"], decision["official_entity_code"]), {}
        )
        rows.append(
            {
                **decision,
                "product_name": products[decision["product_id"]]["canonical_name_ru"],
                "official_entity_name_ru": entity.get("official_entity_name_ru", ""),
                "official_entity_name_kk": entity.get("official_entity_name_kk", ""),
                "official_source_name": source["name"],
                "source_version": source["version"],
                "source_url": source["current_url"],
            }
        )
    return rows


def build_unresolved_queue(
    products: dict[str, dict[str, str]],
    source_id: str,
    entities: list[dict[str, str]],
    crosswalk: list[dict[str, str]],
    legacy_product_ids: set[str],
) -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    name_index: defaultdict[str, list[dict[str, str]]] = defaultdict(list)
    for entity in entities:
        name_index[normalize(entity["official_entity_name_ru"])].append(entity)
    decided = {
        row["product_id"]
        for row in crosswalk
        if row["official_source_id"] == source_id
    }
    unresolved = []
    conflicts = []
    for product_id in sorted(products):
        if product_id in decided:
            continue
        product = products[product_id]
        candidates = sorted(
            name_index.get(normalize(product["canonical_name_ru"]), []),
            key=lambda row: row["official_entity_code"],
        )
        codes = [row["official_entity_code"] for row in candidates]
        names = [row["official_entity_name_ru"] for row in candidates]
        if len(candidates) > 1:
            reason = "AMBIGUOUS_IDENTICAL_NORMALIZED_NAME"
            conflicts.append(
                {
                    "conflict_type": reason,
                    "product_id": product_id,
                    "product_name": product["canonical_name_ru"],
                    "official_source_id": source_id,
                    "official_entity_codes": codes,
                    "official_entity_names": names,
                    "resolution": "MANUAL_REVIEW_REQUIRED",
                }
            )
        elif candidates:
            reason = "IDENTICAL_NAME_CANDIDATE_NOT_SEMANTICALLY_APPROVED"
        else:
            reason = "NOT_RESEARCHED_NO_SAFE_AUTOMATIC_CANDIDATE"
        unresolved.append(
            {
                "product_id": product_id,
                "product_name": product["canonical_name_ru"],
                "product_status": product["status"],
                "official_source_id": source_id,
                "legacy_bns_reference": "YES" if product_id in legacy_product_ids else "NO",
                "candidate_official_codes": codes,
                "candidate_official_names": names,
                "status": "REVIEW_REQUIRED",
                "reason": reason,
            }
        )
    return unresolved, conflicts
