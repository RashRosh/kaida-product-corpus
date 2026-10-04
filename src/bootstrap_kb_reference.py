"""Export the reviewed Excel snapshot into repository-native KB seed files."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "reference" / "KAIDA_Master_KB_Iteration_3.xlsx"
DEFAULT_KB_DIR = ROOT / "kb"
EXPECTED_SHA256 = "74cb60b07c970bd83151dd1cd226b06e11d2798c5d7bf6da53ae31807333213a"
REQUIRED_SHEETS = {
    "PRODUCTS",
    "ALIASES",
    "ATTRIBUTE_DEFINITIONS",
    "RULES",
    "DECISION_LOG",
}

SEED_EXPORTS = {
    "PRODUCTS": "products.csv",
    "ALIASES": "aliases.csv",
    "ATTRIBUTE_DEFINITIONS": "attribute_definitions.csv",
    "MAPPINGS": "reference_mappings.csv",
    "OBSERVED_NAMES": "reference_observed_names.csv",
    "UNRESOLVED": "reference_unresolved_top.csv",
    "RULES": "reference_rules.csv",
    "DECISION_LOG": "reference_decision_log.csv",
}


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def sheet_rows(sheet: Any) -> tuple[list[str], list[dict[str, str]]]:
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        raise ValueError(f"Sheet {sheet.title} is empty")
    headers = [cell_text(value).strip() for value in values[0]]
    if not all(headers):
        raise ValueError(f"Sheet {sheet.title} has blank headers")
    rows: list[dict[str, str]] = []
    for values_row in values[1:]:
        if not any(value is not None and cell_text(value).strip() for value in values_row):
            continue
        rows.append(
            {header: cell_text(value) for header, value in zip(headers, values_row)}
        )
    return headers, rows


def write_csv(path: Path, headers: list[str], rows: Iterable[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def action_for(reference_decision: str) -> str:
    decision = reference_decision.upper()
    if "MALFORMED" in decision:
        return "MALFORMED"
    if "OUT_OF_SCOPE" in decision or "CATEGORY_HEADER" in decision:
        return "OUT_OF_SCOPE"
    if "ALIAS" in decision:
        return "ALIAS"
    if "ATTRIBUTES_ONLY" in decision or decision in {"ATTRIBUTE", "SAME_PRODUCT"}:
        return "ATTRIBUTES_ONLY"
    if "NEW_PRODUCT" in decision:
        return "CREATE_PRODUCT_CANDIDATE"
    if "KEEP_EXISTING" in decision:
        return "MAP_EXISTING"
    return "DEFER"


def decision_records(
    rows: Iterable[dict[str, str]], products: Iterable[dict[str, str]]
) -> list[dict[str, Any]]:
    product_ids_by_name: dict[str, str] = {}
    duplicate_names: set[str] = set()
    for product in products:
        name = product.get("canonical_name_ru", "").strip()
        if not name:
            continue
        if name in product_ids_by_name:
            duplicate_names.add(name)
        product_ids_by_name[name] = product.get("product_id", "")
    for name in duplicate_names:
        product_ids_by_name.pop(name, None)

    records = []
    for row in rows:
        authority = row.get("authority", "")
        canonical_name = row.get("canonical_or_old", "").strip()
        reference_product_id = product_ids_by_name.get(canonical_name, "")
        action = action_for(row.get("decision", ""))
        product_id = (
            reference_product_id
            if action in {"MAP_EXISTING", "ALIAS", "ATTRIBUTES_ONLY"}
            else ""
        )
        product_name = canonical_name if product_id else ""
        rationale = row.get("note", "")
        if action == "MAP_EXISTING" and not product_id:
            action = "DEFER"
            suffix = "Reference MAP_EXISTING had no unique explicit Product target."
            rationale = f"{rationale} {suffix}".strip()
        records.append(
            {
                "decision_id": row.get("decision_id", ""),
                "scope": row.get("subject", ""),
                "category_condition": "",
                "context_condition": "",
                "action": action,
                "product_id": product_id,
                "product_name": product_name,
                "reference_product_id": reference_product_id,
                "reference_product_name": canonical_name,
                "attributes": {},
                "confidence": "HIGH" if authority == "USER" else "PROVISIONAL",
                "provenance": row.get("source", ""),
                "rationale": rationale,
                "created_at": "",
                "version": "iteration-3",
                "authority": authority,
                "reference_decision": row.get("decision", ""),
            }
        )
    return records


def write_jsonl(path: Path, records: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False, sort_keys=True))
            handle.write("\n")


def bootstrap(workbook_path: Path, kb_dir: Path) -> dict[str, int]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Reference workbook is missing: {workbook_path}")
    actual_hash = file_sha256(workbook_path)
    if actual_hash != EXPECTED_SHA256:
        raise ValueError(
            f"Reference workbook SHA-256 mismatch: {actual_hash}; expected {EXPECTED_SHA256}"
        )

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    missing = sorted(REQUIRED_SHEETS.difference(workbook.sheetnames))
    if missing:
        raise ValueError(f"Reference workbook is missing sheets: {', '.join(missing)}")

    extracted: dict[str, tuple[list[str], list[dict[str, str]]]] = {}
    for sheet_name, filename in SEED_EXPORTS.items():
        if sheet_name not in workbook.sheetnames:
            continue
        headers, rows = sheet_rows(workbook[sheet_name])
        extracted[sheet_name] = (headers, rows)
        write_csv(kb_dir / "seed" / filename, headers, rows)

    products = extracted["PRODUCTS"][1]
    decisions = decision_records(extracted["DECISION_LOG"][1], products)
    write_jsonl(kb_dir / "manual_decisions.jsonl", decisions)

    manifest = {
        "workbook": workbook_path.name,
        "sha256": actual_hash,
        "sheets": {
            name: len(rows) for name, (_, rows) in sorted(extracted.items())
        },
    }
    (kb_dir / "seed" / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return {name: len(rows) for name, (_, rows) in extracted.items()}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--kb-dir", type=Path, default=DEFAULT_KB_DIR)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    counts = bootstrap(args.workbook, args.kb_dir)
    print(f"Verified workbook SHA-256: {EXPECTED_SHA256}")
    for sheet_name, count in sorted(counts.items()):
        print(f"{sheet_name}: {count} rows")


if __name__ == "__main__":
    main()
